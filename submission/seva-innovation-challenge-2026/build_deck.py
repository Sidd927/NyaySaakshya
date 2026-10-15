"""
Build the NyaySaakshya Seva Innovation Challenge 2026 deck from the OFFICIAL template.

    /tmp/pptxenv/bin/python build_deck.py            # needs python-pptx; run from this folder

Reads  reference/Seva_Innovation_Challenge_2026_Presentation_Template.pptx  (never modified)
Writes NyaySaakshya_Seva_Innovation_Challenge_2026.pptx

The official template's six slides, header (number, title), footer and background are kept.
Template guidance text is replaced; content is drawn with plain shapes and text in the
template's font (Nirmala UI) and palette. Every statement is mapped to evidence in
docs/submission/PHASE_H_CLAIM_AUDIT.md. The deck carries no placeholders: details only the
applicant can confirm (contact, qualification, mentor) are left out unless filled in
APPLICANT below, and are listed in docs/submission/USER_INPUT_REQUIRED.md.
"""
import os

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(HERE, "reference", "Seva_Innovation_Challenge_2026_Presentation_Template.pptx")
OUT = os.path.join(HERE, "NyaySaakshya_Seva_Innovation_Challenge_2026.pptx")
ASSETS = os.path.join(HERE, "assets")

# Name as shown on the portal (taken from Aadhaar). Check the spelling against the portal.
# Optional details: leave as None to omit the line from the deck. Never guess these.
APPLICANT = {
    "name": "Siddhant Hiralal Patil",
    "contact": None,        # e.g. "name@example.com | +91 ..." (applicant's own, confirmed)
    "qualification": None,  # e.g. "B.E. Computer Engineering, 2025" (applicant's own, confirmed)
    "mentor": None,         # e.g. "Dr A. Kulkarni, Professor, XYZ College" (only if real)
}
SECTOR = "Law and Order"
SUB_THEME = "Cyber-fraud, digital extortion and financial-crime intelligence"

NAVY, ORANGE, SLATE, MUTED = "1B2A5C", "F28C28", "2F3A56", "6B7280"
LINE, WHITE = "D9DEEA", "FFFFFF"
CARD_DARK, RING_DARK, ON_DARK, ON_DARK_DIM = "253870", "3A4C85", "DCE3F5", "C9D2EC"
WARN = "B4531B"
FONT = "Nirmala UI"
A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

prs = Presentation(TEMPLATE)
slides = list(prs.slides)


# ----------------------------------------------------------------------------- helpers
def by_name(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    raise KeyError(name)


def keep_only(slide, names):
    for sh in list(slide.shapes):
        if sh.name not in names:
            sh._element.getparent().remove(sh._element)


def _rgb(hex_):
    return RGBColor.from_string(hex_)


def _fill_runs(p, segs, size, color, bold, italic):
    for txt, ov in segs:
        r = p.add_run()
        r.text = txt
        f = r.font
        f.name = FONT
        f.size = Pt(ov.get("size", size))
        f.bold = ov.get("bold", bold)
        f.italic = ov.get("italic", italic)
        f.color.rgb = _rgb(ov.get("color", color))


def write(tf, paras, size=10, color=SLATE, bold=False, italic=False, align=None, space_after=0):
    """Fill a text frame. A paragraph is a string, a list of (text, overrides) segments, or
    (segments, options) where options may hold bullet / space_after / align."""
    first = True
    for para in paras:
        opts = {}
        if isinstance(para, tuple) and len(para) == 2 and isinstance(para[1], dict) and not isinstance(para[0], tuple):
            para, opts = para
        segs = [(para, {})] if isinstance(para, str) else para
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        a = opts.get("align", align)
        if a is not None:
            p.alignment = a
        p.space_after = Pt(opts.get("space_after", space_after))
        _fill_runs(p, segs, size, color, bold, italic)
        ppr = p._p.get_or_add_pPr()
        for tag in ("buNone", "buChar"):
            for el in ppr.findall(A_NS + tag):
                ppr.remove(el)
        if opts.get("bullet"):
            ppr.set("marL", "128016")
            ppr.set("indent", "-128016")
            etree.SubElement(ppr, A_NS + "buChar").set("char", "•")
        else:
            ppr.set("marL", "0")
            ppr.set("indent", "0")
            etree.SubElement(ppr, A_NS + "buNone")


def text(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, **kw):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    write(tf, paras if isinstance(paras, list) else [paras], **kw)
    return tb


def box(slide, x, y, w, h, fill=WHITE, line=LINE, radius=0.045, lw=0.75, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = _rgb(fill)
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = _rgb(line)
        sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    return sh


def set_header(slide, sub_text, dark=False):
    sub = by_name(slide, "Text 3")
    para = sub.text_frame.paragraphs[0]
    r = para.runs[0]
    r.text = sub_text
    for extra in para.runs[1:]:
        extra._r.getparent().remove(extra._r)
    r.font.size = Pt(11.5)
    r.font.italic = False
    r.font.color.rgb = _rgb(ON_DARK if dark else SLATE)


def notes(slide, body):
    slide.notes_slide.notes_text_frame.text = body


def pic(slide, name, x, y, w):
    return slide.shapes.add_picture(os.path.join(ASSETS, name), Inches(x), Inches(y), width=Inches(w))


def initials_of(name):
    caps = "".join(w[0] for w in name.split() if w[0].isupper())
    return caps[0] + caps[-1] if len(caps) > 2 else caps


# ============================================================================= slide 1
s = slides[0]
keep_only(s, {"Text 0", "Text 1", "Shape 15"})
by_name(s, "Text 0").text_frame.paragraphs[0].runs[0].text = "SEVA INNOVATION CHALLENGE 2026  •  MAHARASHTRA ZONE"
title = by_name(s, "Text 1")
title.text_frame.paragraphs[0].runs[0].text = "NyaySaakshya"
title.left, title.top, title.width, title.height = Inches(0.5), Inches(0.78), Inches(4.6), Inches(0.75)

text(s, 0.5, 1.62, 4.6, 0.7, ["Know what an email-traffic capture proves, and what it cannot."],
     size=17, bold=True, color=WHITE)
text(s, 0.5, 2.42, 4.55, 1.0,
     ["A working prototype for forensic triage of saved email-traffic captures. Each issue it reports cites the "
      "packets behind it. When a capture cannot support a score, NyaySaakshya says so. It does not detect "
      "fraud or identify suspects."], size=10.5, color=ON_DARK)

snap = [("Applicant", [(APPLICANT["name"], {"bold": True, "color": WHITE})], 0.27),
        ("Application", [("Individual (team size 1)", {})], 0.27),
        ("Sector", [(SECTOR, {})], 0.27),
        ("Sub-theme", [(SUB_THEME, {})], 0.43),
        ("Stage", [("Working Prototype", {"bold": True, "color": WHITE}),
                   (": lab-validated, no external pilot", {})], 0.27)]
if APPLICANT["contact"]:
    snap.append(("Contact", [(APPLICANT["contact"], {})], 0.27))
yy = 3.62
for label, val, h in snap:
    text(s, 0.5, yy, 0.95, h, [label], size=9.5, color=ON_DARK_DIM)
    text(s, 1.45, yy, 3.6, h, [val], size=10, color=ON_DARK)
    yy += h

panel = by_name(s, "Shape 15")
panel.left, panel.top, panel.width, panel.height = Inches(5.3), Inches(0.9), Inches(4.2), Inches(4.3)
text(s, 5.48, 1.02, 3.9, 0.25, ["Real output of the current build, synthetic lab captures"],
     size=9.5, bold=True, color=ORANGE)
W1 = 3.84
pic(s, "screen_scored.png", 5.48, 1.34, W1)
h1 = W1 * 456 / 1374
text(s, 5.48, 1.34 + h1 + 0.06, 3.84, 0.4,
     ["Enough evidence: a score, with the rule, standard and packet it rests on"], size=9.5, color=ON_DARK)
y2 = 1.34 + h1 + 0.5
pic(s, "screen_withheld.png", 5.48, y2, W1)
h2 = W1 * 512 / 1374
text(s, 5.48, y2 + h2 + 0.06, 3.84, 0.4,
     ["Not enough evidence: no score, and the reason is stated"], size=9.5, color=ON_DARK)
text(s, 5.48, 4.86, 3.84, 0.25, ["Cropped from the working console; not edited."], size=8.5,
     italic=True, color=ON_DARK_DIM)
notes(s, "Sources and basis. Screens: the NyaySaakshya console running locally on 15 Oct 2026, uploading two synthetic "
         "lab captures (a multi-session SMTP capture with plaintext authentication, and a capture that ends inside a TLS "
         "handshake). Sector and sub-theme are as listed on the Seva Maharashtra-zone portal. The sub-theme names "
         "financial-crime intelligence; NyaySaakshya supports the evidence-handling step and does not itself provide "
         "intelligence or determine that fraud occurred.")

# ============================================================================= slide 2
s = slides[1]
keep_only(s, {"Shape 0", "Text 1", "Text 2", "Text 3", "Text 4"})
set_header(s, "A capture can show less than it seems to. Reading too much into it can undermine the finding built on it.")
text(s, 0.5, 1.26, 6.0, 0.28, ["Three ways a capture can mislead an investigator"], size=12.5, bold=True, color=NAVY)
for x, w, label in ((0.68, 1.75, "What the capture shows"), (2.65, 1.6, "Easy to conclude"),
                    (4.45, 2.0, "What it actually supports")):
    text(s, x, 1.6, w, 0.22, [label], size=9, bold=True, color=MUTED)
rows = [("It ends before the TLS handshake finishes", "No encryption was used",
         "Incomplete: the capture stops before the answer is known"),
        ("The server never offers STARTTLS", "Someone stripped the encryption",
         "Ambiguous: genuine non-support and stripping look identical in the bytes"),
        ("TLS 1.3 encrypts the server certificate", "The certificate is invalid",
         "Not observable in a passive capture")]
ry = 1.86
for a, b, c in rows:
    box(s, 0.5, ry, 6.05, 0.86, fill=WHITE, line=LINE)
    box(s, 0.5, ry, 0.06, 0.86, fill=ORANGE, line=None, shape=MSO_SHAPE.RECTANGLE)
    text(s, 0.68, ry, 1.75, 0.86, [a], size=10, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 2.65, ry, 1.6, 0.86, [b], size=10, color=WARN, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 4.45, ry, 2.0, 0.86, [c], size=10, color=SLATE, anchor=MSO_ANCHOR.MIDDLE)
    ry += 0.94
box(s, 6.8, 1.26, 2.7, 3.38, fill=NAVY, line=NAVY, radius=0.04)
text(s, 6.98, 1.4, 2.35, 3.15, [
    [("Why it matters", {"size": 11.5, "bold": True, "color": ORANGE})],
    [("A conclusion the evidence cannot support is a weak foundation for any cyber-fraud case built on it.", {"size": 10, "color": WHITE})],
    [("Who is affected", {"size": 11.5, "bold": True, "color": ORANGE})],
    [("Cyber-cell and forensic investigators, and the supervisors who must defend a finding. Proposed first "
      "users: a Maharashtra cyber cell or forensic lab. No partner has agreed yet.", {"size": 10, "color": WHITE})],
    [("The gap today", {"size": 11.5, "bold": True, "color": ORANGE})],
    [("Packet tools show packets. Deciding what is safe to conclude is manual and hard to repeat.",
      {"size": 10, "color": WHITE})],
], space_after=4)
box(s, 0.5, 4.74, 9.0, 0.5, fill="E8ECF6", line=None, radius=0.1)
text(s, 0.68, 4.74, 8.65, 0.5,
     [[("Context only: ", {"bold": True}),
       ("PIB (7 Oct 2025) cites over 17.82 lakh complaints in the national CFCFRMS cyber-fraud system; period not "
        "stated. It is not email-specific, and we make no claim about how many cases involve email.", {})]],
     size=9, color=SLATE, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "Sources. Failure modes: RFC 3207 (SMTP STARTTLS, Feb 2002) notes that an active attacker can delete the "
         "'250 STARTTLS' response, and the same bytes arise when a server does not support STARTTLS. RFC 8446 (TLS 1.3) "
         "encrypts the server certificate, so a passive capture cannot show it. Context figure: PIB backgrounder "
         "'Curbing Cyber Frauds in Digital India', 7 Oct 2025: through CFCFRMS financial institutions saved more than "
         "Rs 5,489 crore in over 17.82 lakh complaints; the period is not stated. The share involving email is not "
         "published, so none is claimed.")

# ============================================================================= slide 3
s = slides[2]
keep_only(s, {"Shape 0", "Text 1", "Text 2", "Text 3", "Text 4"})
set_header(s, "Every fact is labelled by how strongly the capture supports it. If it cannot support a score, none is given.")
steps = [("Capture", "Saved PCAP or PCAPNG; SHA-256 recorded"), ("Analysis", "TShark parses; sessions rebuilt"),
         ("Evidence", "Every fact gets a state"), ("Finding", "Rule, standard and frame cited"),
         ("Posture", "A score, or none"), ("Report", "JSON, HTML and PDF")]
sx, sw, gap = 0.5, 1.37, 0.156
for i, (lab, desc) in enumerate(steps):
    x = sx + i * (sw + gap)
    box(s, x, 1.26, sw, 0.72, fill=WHITE, line=LINE)
    text(s, x + 0.08, 1.3, sw - 0.16, 0.25, [lab], size=11, bold=True, color=NAVY)
    text(s, x + 0.08, 1.54, sw - 0.16, 0.42, [desc], size=8.5, color=SLATE)
    if i < 5:
        text(s, x + sw, 1.46, gap, 0.3, ["›"], size=14, bold=True, color=ORANGE, align=PP_ALIGN.CENTER)

text(s, 0.5, 2.08, 4.3, 0.26, ["Six evidence states"], size=12, bold=True, color=NAVY)
states = [("OBSERVED", "Seen directly in the captured bytes", True), ("INFERRED", "Deduced; the basis is stated", True),
          ("UNKNOWN", "Too little evidence to decide", False), ("AMBIGUOUS", "More than one reading fits", False),
          ("INCOMPLETE", "Capture is cut off at that point", False), ("NOT OBSERVABLE", "Cannot be seen passively", False)]
for i, (nm, mean, ok) in enumerate(states):
    col, row = i % 2, i // 2
    x, y = 0.5 + col * 2.2, 2.4 + row * 0.6
    box(s, x, y, 2.1, 0.52, fill=NAVY if ok else WHITE, line=NAVY if ok else ORANGE, lw=1.25, radius=0.08)
    text(s, x + 0.1, y + 0.04, 1.95, 0.2, [nm], size=9.5, bold=True, color=WHITE if ok else WARN)
    text(s, x + 0.1, y + 0.25, 1.95, 0.28, [mean], size=8.5, color=ON_DARK if ok else SLATE)
text(s, 0.5, 4.2, 4.3, 0.2, [[("Navy", {"bold": True, "color": NAVY}), (": can support a conclusion.  ", {}),
                              ("Orange", {"bold": True, "color": WARN}), (": cannot.", {})]], size=9, color=SLATE)

text(s, 5.15, 2.08, 4.35, 0.26, ["One finding, traced to its packets"], size=12, bold=True, color=NAVY)
chain = [("Source", "Capture SHA-256 97b2b61a…; stream #0, 10.0.0.6 to 10.0.0.80:587"),
         ("Evidence", "Frame #7: authentication activity = True (OBSERVED); TLS transition = False (OBSERVED)"),
         ("Rule", "SEC-PLAIN-001, citing NIST SP 800-52 Rev. 2 section 3.1"),
         ("Finding", "Authentication activity without TLS protection: CONFIRMED")]
cy = 2.4
for i, (lab, body) in enumerate(chain):
    h = 0.5 if i == 1 else 0.4
    box(s, 5.15, cy, 4.35, h, fill=WHITE, line=LINE, radius=0.08)
    text(s, 5.27, cy, 0.7, h, [lab], size=9, bold=True, color=ORANGE, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 5.95, cy, 3.45, h, [body], size=9, color=SLATE, anchor=MSO_ANCHOR.MIDDLE)
    cy += h + 0.06
text(s, 5.15, cy, 4.35, 0.2, ["From a real run on a synthetic lab capture."], size=8.5, italic=True, color=MUTED)

box(s, 0.5, 4.6, 5.0, 0.6, fill=NAVY, line=NAVY, radius=0.1)
text(s, 0.65, 4.6, 4.75, 0.6,
     [[("Withholding a score is deliberate. ", {"bold": True, "color": ORANGE}),
       ("If no session can be assessed, or fewer than half yield a conclusion, the result is Not scored, with the "
        "reason. Blindness is never reported as safety.", {})]],
     size=9, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
box(s, 5.62, 4.6, 3.88, 0.6, fill=WHITE, line=LINE, radius=0.1)
text(s, 5.75, 4.6, 3.65, 0.6,
     [[("How it differs: ", {"bold": True, "color": NAVY}),
       ("Wireshark, Zeek and NetworkMiner expose packets and fields. Their documentation, as we read it, has no "
        "per-fact evidence states. A design strength, not a claim of uniqueness.", {})]],
     size=8.5, color=SLATE, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "Sources. Evidence states and abstention are properties of the prototype, shown here from a real local run on a "
         "synthetic capture. NIST SP 800-52 Rev. 2 (Guidelines for the Selection, Configuration, and Use of TLS "
         "Implementations). Comparison basis, documentation read 15 Oct 2026: Wireshark user guide (Expert Information) "
         "and TLS wiki; Zeek SMTP::Info (a tls true/false field); NetworkMiner product page (extracts files, emails and "
         "credentials from PCAPs; not designed to decrypt). No benchmark was run against these tools and no speed or "
         "accuracy comparison is made.")

# ============================================================================= slide 4
s = slides[3]
keep_only(s, {"Shape 0", "Text 1", "Text 2", "Text 3", "Text 4", "Text 5", "Shape 6", "Shape 7", "Text 8", "Shape 9",
              "Text 10", "Shape 11", "Text 12", "Shape 13", "Text 14", "Shape 15", "Text 16"})
set_header(s, "A working prototype, tested internally, and an honest list of what a pilot would still need.")
r = by_name(s, "Text 5").text_frame.paragraphs[0].runs[0]
r.text = "Current stage: Working Prototype (lab-validated; no external pilot)"
r.font.size = Pt(11.5)
r.font.bold = True
r.font.italic = False
r.font.color.rgb = _rgb(NAVY)
tick = by_name(s, "Shape 11")
tick.fill.solid()
tick.fill.fore_color.rgb = _rgb(ORANGE)
tick.text_frame.text = "✓"
tick.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
for rr in tick.text_frame.paragraphs[0].runs:
    rr.font.size = Pt(12)
    rr.font.bold = True
    rr.font.name = FONT
    rr.font.color.rgb = _rgb(WHITE)
wp = by_name(s, "Text 12").text_frame.paragraphs[0].runs[0]
wp.font.bold = True
wp.font.color.rgb = _rgb(ORANGE)

cards = [
    (0.5, 2.95, "Built and tested (internal)", [
        "Python and TShark analysis, React console, FastAPI service, SQLite",
        "1,521 automated tests",
        "98 controlled lab captures: 10 from Postfix and Dovecot software in a lab, 6 OpenSSL handshakes, 82 crafted in-project",
        "Same capture and software give the same result (98 of 98)",
        "Validation found and fixed cases where the engine over-read evidence"],
     "Lab validation only: no real-world traffic and no accuracy benchmark."),
    (3.55, 2.95, "Not built yet", [
        "Login and roles",
        "Audit log",
        "Encrypted storage, backups, retention",
        "Fix for a database-concurrency defect",
        "Enforced analysis time limit",
        "Sandboxed parsing of untrusted files"],
     "Not production-ready. The demo is local, with synthetic captures only."),
    (6.6, 2.9, "Proposed path (not committed)", [
        [("0 to 3 months: ", {"bold": True}), ("close the gaps; write a threat model", {})],
        [("3 to 6 months: ", {"bold": True}), ("technical pilot on synthetic or consented captures", {})],
        [("6 to 12 months: ", {"bold": True}), ("supervised single-unit pilot, only if a department agrees", {})]],
     "Depends on security mentorship, a willing unit and that unit's legal approval for real captures. "
     "Costs not yet estimated."),
]
for x, w, head, items, foot in cards:
    box(s, x, 2.6, w, 2.55, fill=WHITE, line=LINE)
    text(s, x + 0.14, 2.7, w - 0.28, 0.26, [head], size=11.5, bold=True, color=NAVY)
    paras = [(it, {"bullet": True, "space_after": 2}) for it in items]
    text(s, x + 0.14, 3.0, w - 0.28, 1.4, paras, size=9, color=SLATE)
    text(s, x + 0.14, 4.55, w - 0.28, 0.55, [foot], size=8.5, italic=True, color=WARN, anchor=MSO_ANCHOR.BOTTOM)
notes(s, "Basis. Internal validation record: 98 controlled captures (10 sessions from real Postfix and Dovecot software run in a lab, 6 OpenSSL "
         "handshakes, 82 crafted by in-project generators; no real-world traffic); 1,521 automated tests; identical results on repeat runs for 98 "
         "of 98 captures. No real-world or third-party traffic has been evaluated and no false-positive or false-negative "
         "rates are claimed. The security gaps listed are the known blockers for any pilot. Timeline items are proposals, "
         "not commitments. No hosting or implementation cost is stated because none has been verified.")

# ============================================================================= slide 5
s = slides[4]
keep_only(s, {"Shape 0", "Text 1", "Text 2", "Text 3", "Text 4"})
set_header(s, "Fewer unsupported conclusions in cyber-fraud cases: the expected benefit, to be measured in a supervised pilot.")
cols = [
    (0.5, "Public value (expected, not yet measured)", [
        "Investigators and supervisors can see what a capture supports and which packets show it",
        "Findings can be re-checked: the same capture gives the same result",
        "Fewer conclusions that the evidence cannot bear",
        "Fits SDG 16, target 16.a: strengthening institutions that combat crime"], None),
    (3.55, "Adoption path (proposed)", None,
     [("1", "Hardened local build"), ("2", "One supervised unit"), ("3", "Measured results, then decide on more units")]),
    (6.6, "How success would be measured", [
        "Time to a first assessment",
        "Time to a finished report",
        "Findings with usable packet citations",
        "Handling of unresolved evidence",
        "Repeat-run agreement",
        "Supervisor confidence"], None),
]
for x, head, bullets, flow in cols:
    box(s, x, 1.28, 2.9, 2.62, fill=WHITE, line=LINE)
    text(s, x + 0.14, 1.38, 2.62, 0.4, [head], size=11, bold=True, color=NAVY)
    if bullets:
        text(s, x + 0.14, 1.82, 2.62, 1.5, [(b, {"bullet": True, "space_after": 3}) for b in bullets], size=9.5, color=SLATE)
    if flow:
        fy = 1.85
        for num, lab in flow:
            box(s, x + 0.14, fy, 2.62, 0.42, fill="EEF1F8", line=None, radius=0.15)
            box(s, x + 0.2, fy + 0.07, 0.28, 0.28, fill=ORANGE, line=None, shape=MSO_SHAPE.OVAL)
            text(s, x + 0.2, fy + 0.07, 0.28, 0.28, [num], size=9.5, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
                 anchor=MSO_ANCHOR.MIDDLE)
            text(s, x + 0.56, fy, 2.15, 0.42, [lab], size=9.5, color=SLATE, anchor=MSO_ANCHOR.MIDDLE)
            fy += 0.5
        text(s, x + 0.14, fy + 0.02, 2.62, 0.5,
             ["Each unit would run it locally, so captures stay in its own environment. A design intention, not built or validated."],
             size=8.5, italic=True, color=MUTED)
text(s, 6.74, 3.35, 2.62, 0.4, ["Baselines come from the pilot. No targets are claimed yet."], size=8.5, italic=True, color=MUTED)

box(s, 0.5, 4.05, 5.0, 1.05, fill=NAVY, line=NAVY, radius=0.06)
text(s, 0.68, 4.15, 4.65, 0.9, [
    [("What the Challenge would unlock", {"size": 11, "bold": True, "color": ORANGE})],
    [("Security mentorship and one supervised pilot would turn a lab-validated prototype into a measured "
      "public-sector trial of whether it helps investigators avoid unsupported conclusions.", {"size": 10, "color": WHITE})],
], space_after=3)
box(s, 5.62, 4.05, 3.88, 1.05, fill=WHITE, line=LINE, radius=0.06)
text(s, 5.78, 4.13, 3.6, 0.95, [
    [("Not claimed: ", {"bold": True, "color": WARN}),
     ("beneficiaries reached, time or cost saved, jobs. None has been measured.", {})],
    [("Sustainability: ", {"bold": True, "color": NAVY}),
     ("no revenue model is decided. It builds on open-source components (TShark, Python, FastAPI, React); "
      "NyaySaakshya's own licence is not yet chosen.", {})],
], size=8.5, color=SLATE, space_after=3)
notes(s, "Basis. No beneficiary, saving or jobs figure is claimed: there is no measured baseline and no pilot. SDG 16 target "
         "16.a (United Nations): strengthen relevant national institutions to prevent violence and combat terrorism and "
         "crime. The adoption path and the local-processing model are proposals. Third-party components: TShark/Wireshark "
         "(GPL-2.0-or-later, used as a separate program), FastAPI and React (MIT), ReportLab (BSD).")

# ============================================================================= slide 6
s = slides[5]
keep_only(s, {"Shape 0", "Text 1", "Text 2", "Text 3", "Text 43"})
set_header(s, "One applicant with a working prototype, and the specific help needed to reach a supervised pilot.", dark=True)
box(s, 0.5, 1.28, 3.75, 1.25, fill=CARD_DARK, line=CARD_DARK, radius=0.07)
box(s, 0.68, 1.5, 0.8, 0.8, fill=RING_DARK, line=ORANGE, lw=1.5, shape=MSO_SHAPE.OVAL)
text(s, 0.68, 1.5, 0.8, 0.8, [initials_of(APPLICANT["name"])], size=16, bold=True, color=WHITE,
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, 1.65, 1.46, 2.5, 0.9, [
    [(APPLICANT["name"], {"size": 13.5, "bold": True, "color": WHITE})],
    [("Applicant and team lead", {"size": 10, "color": ORANGE})],
    [(APPLICANT["qualification"] or "Individual application (team size 1)", {"size": 9.5, "color": ON_DARK})],
], space_after=3)
box(s, 0.5, 2.68, 3.75, 2.0 if not APPLICANT["mentor"] else 1.7, fill=CARD_DARK, line=CARD_DARK, radius=0.05)
text(s, 0.68, 2.78, 3.4, 1.85, [
    [("Next milestone (proposed, about 3 months)", {"size": 10.5, "bold": True, "color": ORANGE})],
    [("A hardened local pilot-candidate build: login and roles, audit log, encrypted storage, backups and a written "
      "threat model.", {"size": 10, "color": WHITE})],
    [("Why this help", {"size": 10.5, "bold": True, "color": ORANGE})],
    [("The core analysis is built. What blocks a pilot is security engineering and access to a supervised unit.",
      {"size": 10, "color": WHITE})],
    [("Individual application: no other team member, mentor, partner or funder is named.",
      {"size": 9, "color": ON_DARK_DIM})],
], space_after=5)
if APPLICANT["mentor"]:
    text(s, 0.5, 4.45, 3.75, 0.3, ["Mentor: " + APPLICANT["mentor"]], size=9.5, color=ON_DARK)
if APPLICANT["contact"]:
    text(s, 0.5, 4.8, 3.75, 0.3, ["Contact: " + APPLICANT["contact"]], size=9.5, color=ON_DARK)

box(s, 4.45, 1.28, 5.05, 3.82, fill=WHITE, line=WHITE, radius=0.04)
text(s, 4.65, 1.4, 4.7, 0.28, ["Support we are seeking"], size=12.5, bold=True, color=NAVY)
asks = [(True, "Security mentorship", "Review the threat model and hardening plan for sensitive evidence"),
        (True, "A supervised pilot unit", "Cyber cell or forensic lab; synthetic data first. No partner yet"),
        (True, "Testing, certification and IP guidance", "Choose a licence; understand evidence-handling expectations"),
        (True, "Funding for a secure pilot build", "Login, audit log, encrypted storage and backups"),
        (False, "Incubation and workspace", "Not requested")]
ay = 1.8
for ok, head, sub in asks:
    box(s, 4.68, ay + 0.03, 0.22, 0.22, fill=ORANGE if ok else WHITE, line=ORANGE, lw=1.25, radius=0.1)
    if ok:
        text(s, 4.68, ay + 0.03, 0.22, 0.22, ["✓"], size=9, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE)
    text(s, 5.05, ay, 4.3, 0.55, [[(head, {"size": 10.5, "bold": ok, "color": NAVY if ok else MUTED})],
                                  [(sub, {"size": 9, "color": SLATE if ok else MUTED})]], space_after=1)
    ay += 0.58
text(s, 4.65, 4.72, 4.7, 0.36, ["Thank you  |  " + APPLICANT["name"]], size=10.5, bold=True, color=NAVY)
notes(s, "Individual application, team size 1; name as shown on the portal. No other team member, mentor, partner or funder "
         "is named. Support ticks are the applicant's request; incubation is not requested.")

prs.save(OUT)
print("saved", OUT, "slides:", len(prs.slides))
